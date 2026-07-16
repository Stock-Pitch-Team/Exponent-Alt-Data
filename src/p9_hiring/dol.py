"""Stream DOL H-1B LCA + PERM disclosure workbooks, keep only Exponent rows.

Files are ~70-140MB each with ~100k-500k rows; openpyxl read_only streams them
row-wise. Filtered rows are cached per file so re-runs are instant.
"""
import json
import logging
import re

from openpyxl import load_workbook

from config import settings

log = logging.getLogger(__name__)

DOL_DIR = settings.RAW_DIR / "dol"
CACHE_DIR = settings.INTERIM_DIR / "dol_filtered"


class UnknownSchema(Exception):
    """A workbook's header row has no recognisable employer column."""

_EMPLOYER_RE = re.compile(r"\bEXPONENT\b", re.IGNORECASE)
# Columns we keep, matched loosely against header names. DOL rewrote the PERM
# schema at FY2025-Q3 (EMPLOYER_NAME -> EMP_BUSINESS_NAME, PW_* -> PWD_*,
# JOB_INFO_WORK_* -> PRIMARY_WORKSITE_*); both generations are listed here.
_WANTED = {
    "employer": ("EMPLOYER_NAME", "EMP_BUSINESS_NAME"),
    "status": ("CASE_STATUS",),
    "received": ("RECEIVED_DATE", "CASE_RECEIVED_DATE"),
    "decision": ("DECISION_DATE",),
    "title": ("JOB_TITLE",),
    "soc_title": ("SOC_TITLE", "PW_SOC_TITLE", "PWD_SOC_TITLE"),
    "soc_code": ("SOC_CODE", "PW_SOC_CODE", "PWD_SOC_CODE"),
    "city": ("WORKSITE_CITY", "WORKSITE_CITY_1", "PW_WORKSITE_CITY", "JOB_INFO_WORK_CITY",
             "PRIMARY_WORKSITE_CITY"),
    "state": ("WORKSITE_STATE", "WORKSITE_STATE_1", "PW_WORKSITE_STATE", "JOB_INFO_WORK_STATE",
              "PRIMARY_WORKSITE_STATE"),
    "wage": ("WAGE_RATE_OF_PAY_FROM", "WAGE_RATE_OF_PAY_FROM_1", "PW_AMOUNT_9089",
             "WAGE_OFFER_FROM_9089", "WAGE_OFFERED_FROM_9089", "JOB_OPP_WAGE_FROM"),
    "wage_level": ("PW_WAGE_LEVEL", "PW_WAGE_LEVEL_1", "PW_SKILL_LEVEL", "PW_LEVEL_9089"),
    "employment_start": ("BEGIN_DATE", "EMPLOYMENT_START_DATE", "JOB_INFO_START_DATE",
                         "RECR_INFO_JOB_START_DATE"),
}


def _header_map(header_row) -> dict:
    idx = {}
    headers = [str(c).strip().upper() if c is not None else "" for c in header_row]
    for key, options in _WANTED.items():
        for opt in options:
            if opt in headers:
                idx[key] = headers.index(opt)
                break
    return idx


def filter_file(path) -> list[dict]:
    """Return Exponent rows from one workbook (cached)."""
    CACHE_DIR.mkdir(parents=True, exist_ok=True)
    cache = CACHE_DIR / (path.stem + ".json")
    if cache.exists():
        return json.loads(cache.read_text(encoding="utf-8"))
    log.info("streaming %s ...", path.name)
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows_iter = ws.iter_rows(values_only=True)
    header = next(rows_iter)
    idx = _header_map(header)
    if "employer" not in idx:
        # Never cache/return [] here: an unrecognised schema is a BUG, and a silent
        # empty list is indistinguishable from "Exponent filed nothing this quarter".
        # DOL rewrote the PERM schema at FY2025-Q3 and this path hid four files.
        wb.close()
        raise UnknownSchema(
            f"{path.name}: no employer column among {_WANTED['employer']}; "
            f"headers start with {[str(c) for c in header[:8]]}")
    out, scanned = [], 0
    for row in rows_iter:
        scanned += 1
        emp = row[idx["employer"]]
        if not emp or not _EMPLOYER_RE.search(str(emp)):
            continue
        rec = {"file": path.name, "program": "PERM" if "PERM" in path.name else "LCA",
               "employer": str(emp)}
        for key, i in idx.items():
            if key == "employer":
                continue
            val = row[i]
            rec[key] = str(val)[:80] if val is not None else None
        out.append(rec)
    wb.close()
    cache.write_text(json.dumps(out, indent=0), encoding="utf-8")
    log.info("%s: %d Exponent rows of %d scanned", path.name, len(out), scanned)
    return out


def all_rows() -> list[dict]:
    rows = []
    for path in sorted(DOL_DIR.glob("*.xlsx")):
        rows.extend(filter_file(path))
    return rows
